"""Extract source-native policy/claims panels after verified workbook audit.

NOT causal inference, causal-ready data, or a main warehouse merge.
Keep California's regional claims and both exposure measures separate.
"""
import csv
import datetime as dt
import hashlib
import io
import json
import math
from collections import defaultdict
from pathlib import Path
from urllib.parse import urlparse
from urllib.request import Request, urlopen
from openpyxl import load_workbook

OUT = Path('data')
TARIFF = 'https://raw.githubusercontent.com/dkaenzig/narrative-tariff-shocks/main/narrative-tariff-shocks.xlsx'
CAIT = 'https://capolicylab.org/data/cait-monthly-claims.xlsx'
HEADERS = ['country_iso3','year','period','series','value','unit','frequency','provider',
           'source_url','disaggregation','observation_kind','source_status']
US = 'USA'

def get_xlsx(url, allowed):
    with urlopen(Request(url,headers={'User-Agent':'EconomicPolicyResearch/1.0',
                                      'Accept':'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet,*/*'}),
                 timeout=65) as resp:
        final=resp.geturl()
        if urlparse(final).scheme!='https' or urlparse(final).hostname not in allowed:
            raise ValueError('Unexpected source redirect: '+final)
        raw=resp.read(14_000_001)
    if len(raw)>14_000_000 or not raw.startswith(b'PK'):
        raise ValueError('Unexpected XLSX response')
    return load_workbook(io.BytesIO(raw),read_only=True,data_only=True),hashlib.sha256(raw).hexdigest(),len(raw)

def write_csv(path,records):
    if not records: raise ValueError('Cannot write empty panel: '+str(path))
    path.parent.mkdir(parents=True,exist_ok=True)
    with path.open('w',newline='',encoding='utf-8') as file:
        w=csv.DictWriter(file,fieldnames=HEADERS,extrasaction='raise')
        w.writeheader();w.writerows(records)
    return {'path':str(path),'observations':len(records),'sha256':hashlib.sha256(path.read_bytes()).hexdigest()}

def row(year, period, series, value, unit, freq, provider, source_url, disagg, kind, status):
    val=float(value)
    if not math.isfinite(val):raise ValueError('Nonfinite numeric data')
    return dict(country_iso3=US,year=year,period=period,series=series,value=repr(val),
                unit=unit,frequency=freq,provider=provider,source_url=source_url,
                disaggregation=json.dumps(disagg,ensure_ascii=False,sort_keys=True),
                observation_kind=kind,source_status=status)

def tariff():
    book,sha,nbytes=get_xlsx(TARIFF,{'raw.githubusercontent.com'})
    if book.sheetnames!=['narrative_shock_series']:
        raise ValueError('Unexpected tariff workbook sheets '+str(book.sheetnames))
    sheet=book.active
    records=[];seen=set(); zeros=0;positive=0;negative=0
    for excel_row,values in enumerate(sheet.iter_rows(values_only=True),1):
        if excel_row==1:
            if tuple(values[:2])!=('year','narrative_instrument_distributed'):
                raise ValueError('Tariff column definitions changed')
            continue
        if len(values)<2 or values[0] is None:continue
        y=int(values[0]);v=float(values[1])
        if y in seen or not 1840<=y<=2024:raise ValueError('Duplicate/out-of-range tariff year')
        if not math.isfinite(v):raise ValueError('Bad tariff value')
        seen.add(y);zeros+=int(v==0);positive+=int(v>0);negative+=int(v<0)
        records.append(row(y,str(y),'narrative_instrument_distributed',v,
          'source_native_instrument_codebook_review_required','annual','Den Besten et al (2026)',TARIFF,
          {'source_sheet':'narrative_shock_series','source_excel_row':excel_row},
          'narrative_policy_instrument','SOURCE_NATIVE_NOT_CAUSAL_VALIDATED'))
    if len(records)!=185 or seen!=set(range(1840,2025)) or positive==0 or negative==0:
        raise ValueError('Tariff coverage or sign validation failed')
    file=write_csv(OUT/'tariff'/'processed'/'tariff_shocks_panel.csv',records)
    return dict(dataset='narrative_tariff_shocks_2026',original_url=TARIFF,source_sha256=sha,
       source_bytes=nbytes,year_min=1840,year_max=2024,zeros=zeros,positive=positive,
       negative=negative,validation='SOURCE_SCHEMA_AND_VALUES_CHECKED',research_ready=False,
       warehouse_integrated=False,causal_estimates_computed=False,errors=[],conflicts=[],panel=file,
       notes=['The source instrument is not an observed tariff rate.',
              'Source-native sign and units must be checked against the paper before econometric use.',
              'Zero shock values are preserved, not recoded as missing.','CC BY 4.0, cite original authors.'])

def month(raw,year):
    # Excel stores Stata-labeled dates as Excel serial day numbers in this XLSX.
    # Never interpret e.g. 42736 with Stata epoch (1960-01-01).
    if isinstance(raw,(dt.datetime,dt.date)):
        day=raw.date() if isinstance(raw,dt.datetime) else raw
    elif isinstance(raw,(int,float)):
        day=(dt.datetime(1899,12,30)+dt.timedelta(days=float(raw))).date()
    else: raise ValueError('Unrecognized CAIT month encoding: '+repr(raw))
    if day.day!=1 or day.year!=int(year) or not 2017<=day.year<=2032:
        raise ValueError('CAIT month encoding/period mismatch '+repr((raw,year,day)))
    return day.strftime('%Y-%m')

def parse_rows(sheet,expected):
    iterator=sheet.iter_rows(values_only=True)
    header=next(iterator)
    if tuple(header[:len(expected)])!=expected:
        raise ValueError('Unexpected '+sheet.title+' columns '+repr(header))
    for num,values in enumerate(iterator,2):
        if not values or all(v is None for v in values):continue
        yield num,dict(zip(expected,values[:len(expected)]))

def cait():
    book,sha,nbytes=get_xlsx(CAIT,{'capolicylab.org','www.capolicylab.org'})
    needed={'README','codebook','monthly_claims','monthly_AIscore'}
    if not needed.issubset(book.sheetnames):
        raise ValueError('CAIT required sheets absent: '+repr(book.sheetnames))
    readme=str(book['README']['A1'].value)
    codebook={str(r[0]).strip():str(r[1]) for r in book['codebook'].iter_rows(min_row=2,values_only=True)
              if r[0] and r[1]}
    for key in ('measure','month','q_grp','claims','claims_share','score'):
        if key not in codebook:raise ValueError('CAIT codebook missing '+key)
    if 'California AI Exposure' not in readme:
        raise ValueError('CAIT README identity changed')
    claims_cols=('measure','year','month','q_grp','claims','claims_share')
    score_cols=('measure','year','month','claims','score')
    records=[]; claim_groups=defaultdict(dict); group_count=0; score_count=0
    for excel_row,x in parse_rows(book['monthly_claims'],claims_cols):
        measure=str(x['measure']).strip();yr=int(x['year']);period=month(x['month'],yr)
        group=str(x['q_grp']).strip();key=(measure,period)
        if measure not in {'Anthropic','Eloundou'} or group not in {
             'Top 25%','Bottom 25%','25th–75th percentile'}:
            raise ValueError('Unexpected CAIT measure/exposure group')
        if group in claim_groups[key]:raise ValueError('Duplicate CAIT exposure group')
        if x['claims'] is None or x['claims_share'] is None:
            raise ValueError('Suppressed source cell in base claims tab: missing not zero')
        n=int(x['claims']);share=float(x['claims_share'])
        if n<0 or not 0<=share<=1 or not math.isfinite(share):raise ValueError('Bad claims/share')
        claim_groups[key][group]=(n,share)
        dis={'state':'California','exposure_measure':measure,'exposure_quantile':group,
             'source_sheet':'monthly_claims','source_excel_row':excel_row}
        records.extend([
          row(yr,period,'CAIT_INITIAL_UI_CLAIMS_BY_AI_EXPOSURE',n,'claims','monthly',
              'CA Employment Development Department and California Policy Lab',CAIT,dis,
              'administrative_claim_count','AI_EXPOSURE_NOT_CAUSAL_ATTRIBUTION'),
          row(yr,period,'CAIT_SHARE_OF_CLAIMS_BY_AI_EXPOSURE',share,'fraction_0_1','monthly',
              'CA Employment Development Department and California Policy Lab',CAIT,dis,
              'administrative_claim_share','AI_EXPOSURE_NOT_CAUSAL_ATTRIBUTION')])
        group_count+=1
    for key,groups in claim_groups.items():
        if set(groups)!={'Top 25%','Bottom 25%','25th–75th percentile'}:
            raise ValueError('CAIT incomplete exposure group for '+str(key))
        if abs(sum(v[1] for v in groups.values())-1)>1e-4:
            raise ValueError('CAIT exposure shares fail to sum to one')
    observed_scores=set()
    for excel_row,x in parse_rows(book['monthly_AIscore'],score_cols):
        measure=str(x['measure']).strip();yr=int(x['year']);period=month(x['month'],yr)
        key=(measure,period)
        if measure not in {'Anthropic','Eloundou'} or key in observed_scores or key not in claim_groups:
            raise ValueError('Unexpected, duplicate, or unmatched AI score row')
        observed_scores.add(key)
        if x['claims'] is None or x['score'] is None:
            raise ValueError('Suppressed base AI score cell')
        count=int(x['claims']); score=float(x['score'])
        if count<0 or not math.isfinite(score):raise ValueError('Invalid CAIT score/count')
        if count!=sum(v[0] for v in claim_groups[key].values()):
            raise ValueError('CAIT claim-count totals disagree across tabs')
        dis={'state':'California','exposure_measure':measure,'source_sheet':'monthly_AIscore',
             'source_excel_row':excel_row}
        records.extend([
          row(yr,period,'CAIT_MONTHLY_MEAN_EXPOSURE_SCORE',score,'source_specific_AI_score',
              'monthly','CA Employment Development Department and California Policy Lab',CAIT,dis,
              'mean_exposure_among_UI_claimants','AI_EXPOSURE_NOT_CAUSAL_ATTRIBUTION'),
          row(yr,period,'CAIT_MONTHLY_TOTAL_INITIAL_CLAIMS',count,'claims','monthly',
              'CA Employment Development Department and California Policy Lab',CAIT,dis,
              'administrative_claim_count','DUPLICATED_ACROSS_AI_MEASURES_DO_NOT_SUM')])
        score_count+=1
    if observed_scores!=set(claim_groups) or len(claim_groups)<200:
        raise ValueError('CAIT base claim-score coverage mismatch or too small')
    file=write_csv(OUT/'cait'/'processed'/'cait_panel.csv',records)
    return dict(dataset='cait_ai_exposure_ui_claims_2026',original_url=CAIT,
      source_sha256=sha,source_bytes=nbytes,source_readme=readme,
      source_sheets=list(book.sheetnames),raw_claim_quartile_rows=group_count,
      raw_ai_score_rows=score_count,distinct_measure_months=len(claim_groups),
      measures=sorted({key[0] for key in claim_groups}),validation='SOURCE_SCHEMA_AND_CROSS_TAB_TOTALS_CHECKED',
      research_ready=False,warehouse_integrated=False,causal_estimates_computed=False,
      errors=[],conflicts=[],panel=file,
      notes=['California only; unemployment insurance claims are not all layoffs.',
             'Exposure measure is not observed AI-caused displacement.',
             'Claims appear under both exposure measures: do not sum across Anthropic and Eloundou.',
             'Month dates are parsed as Excel serial dates and checked against source year.',
             'Monthly source figures can be revised; published source reports suppress small cells.',
             'Only the base claims and AIscore tabs are normalized in this pass; 12 detailed disaggregation tabs remain source-only.',
             'Do not republish original CAIT workbook without terms review.'])

def main():
    OUT.mkdir(parents=True,exist_ok=True)
    summaries=[]
    for fun,name in [(tariff,'tariff'),(cait,'cait')]:
        info=fun();info['observations']=info['panel']['observations'];summaries.append(info)
        path=OUT/name/'run_report.json'
        path.write_text(json.dumps(info,indent=2,ensure_ascii=False)+'\n',encoding='utf-8')
    print(json.dumps([{'dataset':x['dataset'],'observations':x['panel']['observations'],
                       'validation':x['validation'],'warehouse_integrated':x['warehouse_integrated']}
                      for x in summaries],indent=2))

if __name__=='__main__':main()
